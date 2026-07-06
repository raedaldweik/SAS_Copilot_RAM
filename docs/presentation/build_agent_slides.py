# Three orchestration-style slides — one per remaining agent — in the same
# visual language as the "SAS Viya Co-Pilot · Agent Orchestration" slide.
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

NAVY  = "0B2242"   # matches the dark pill/bar in the original slide
INK   = "1F2A37"
GRAY  = "6B7280"
FAINT = "9AA5B1"
WHITE = "FFFFFF"
STRIP = "E9EDF4"
PANEL = "D9DFE8"

W, H = 13.333, 7.5

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]


def box(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    return tf

def para(tf, first=False, align=PP_ALIGN.LEFT, space_after=0, line=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    if line:
        p.line_spacing = line
    return p

def run(p, text, size, color, bold=False, italic=False):
    r = p.add_run()
    r.text = text
    f = r.font
    f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, "Arial"
    f.color.rgb = RGBColor.from_string(color)
    return r

def shape(s, kind, x, y, w, h, fill, line=None, adj=None):
    sp = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if adj is not None:
        try:
            sp.adjustments[0] = adj
        except Exception:
            pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = RGBColor.from_string(line)
        sp.line.width = Pt(1.1)
    sp.shadow.inherit = False
    return sp

def arrow(s, kind, x, y, w, h):
    sp = shape(s, kind, x, y, w, h, "C9CFD9")
    return sp

def footer(s):
    tf = box(s, 0.55, H - 0.62, 6.0, 0.5)
    p = para(tf, True, space_after=1)
    run(p, "Company Confidential — For Internal Use Only", 7.5, FAINT)
    p = para(tf)
    run(p, "Copyright © SAS Institute Inc. All rights reserved.", 7.5, FAINT)


def agent_slide(pill_text, subtitle, bar_title, bar_sub, cards, ramp,
                strip_title, strip_sub):
    s = prs.slides.add_slide(BLANK)
    shape(s, MSO_SHAPE.RECTANGLE, 0, 0, W, H, WHITE)
    # content panel outline
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.42, 0.5, W - 0.84, 6.15, WHITE,
          line=PANEL, adj=0.035)

    # title pill
    pw = 0.16 * len(pill_text) / 1.6 + 2.2
    pw = max(5.4, min(9.0, pw))
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, (W - pw) / 2, 0.78, pw, 0.62,
          NAVY, adj=0.5)
    tf = box(s, (W - pw) / 2, 0.78, pw, 0.62, anchor=MSO_ANCHOR.MIDDLE)
    run(para(tf, True, align=PP_ALIGN.CENTER), pill_text, 15, WHITE, bold=True)

    # subtitle
    tf = box(s, 1.5, 1.52, W - 3.0, 0.35)
    run(para(tf, True, align=PP_ALIGN.CENTER), subtitle, 13, GRAY, bold=True)

    # agent bar
    bx, bw = 3.1, W - 6.2
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, bx, 1.98, bw, 0.9, NAVY, adj=0.16)
    tf = box(s, bx + 0.3, 1.98, bw - 0.6, 0.9, anchor=MSO_ANCHOR.MIDDLE)
    p = para(tf, True, align=PP_ALIGN.CENTER, space_after=2)
    run(p, bar_title, 17, WHITE, bold=True)
    p = para(tf, align=PP_ALIGN.CENTER)
    run(p, bar_sub, 10.5, "AFC2DA")

    # card row
    n = len(cards)
    gap = 0.34
    cw = (W - 2 * 0.95 - (n - 1) * gap) / n
    cy, ch = 3.42, 1.95
    for i, (t, d) in enumerate(cards):
        x = 0.95 + i * (cw + gap)
        # down arrow from the agent bar
        arrow(s, MSO_SHAPE.DOWN_ARROW, x + cw / 2 - 0.09, 3.02, 0.18, 0.3)
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, cy, cw, ch, ramp[i],
              adj=0.09)
        tf = box(s, x + 0.18, cy + 0.22, cw - 0.36, 0.62,
                 anchor=MSO_ANCHOR.MIDDLE)
        run(para(tf, True, align=PP_ALIGN.CENTER, line=1.02), t, 12.5, WHITE,
            bold=True)
        tf = box(s, x + 0.2, cy + 0.88, cw - 0.4, ch - 1.0)
        run(para(tf, True, align=PP_ALIGN.CENTER, line=1.12), d, 9.5,
            "F2F6FB")
        if i < n - 1:
            ax = x + cw + gap / 2 - 0.07
            arrow(s, MSO_SHAPE.CHEVRON, ax, cy + ch / 2 - 0.09, 0.16, 0.18)
        # up arrow from the platform strip
        arrow(s, MSO_SHAPE.UP_ARROW, x + cw / 2 - 0.09, cy + ch + 0.12,
              0.18, 0.3)

    # platform strip
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.95, 5.92, W - 1.9, 0.62, STRIP,
          adj=0.18)
    tf = box(s, 1.2, 5.92, W - 2.4, 0.62, anchor=MSO_ANCHOR.MIDDLE)
    p = para(tf, True, align=PP_ALIGN.CENTER, space_after=1)
    run(p, strip_title, 11.5, NAVY, bold=True)
    p = para(tf, align=PP_ALIGN.CENTER)
    run(p, strip_sub, 8.5, GRAY)

    footer(s)
    return s


# ── 1 · Investigation Assistant (teal) ──────────────────────────────
agent_slide(
    "INVESTIGATION ASSISTANT  ·  ALERT TRIAGE",
    "From a queue of alerts to a plan of action — before the first coffee",
    "Investigation Assistant",
    "triage copilot for SAS Visual Investigator — reads, explains, prioritizes",
    [
        ("Prioritized Queue",
         "What to open first —\nand why it matters today."),
        ("Alert, Explained",
         "Why it fired, in plain\nlanguage an auditor can read."),
        ("Network, Mapped",
         "People, companies, accounts —\nand how they connect."),
        ("Action, Recommended",
         "Probable false positive or\nescalate — with next steps."),
    ],
    ["3AA0B5", "1E87A0", "0F6E86", "0B566B"],
    "SAS Visual Investigator  ·  Live Alerts, Entities & Networks",
    "The agent works the same queue your investigators see — every call it makes is visible in the trace",
)

# ── 2 · Procurement Integrity Analyst (green) ───────────────────────
agent_slide(
    "PROCUREMENT INTEGRITY ANALYST  ·  RED FLAGS BEFORE AWARD",
    "Ready-made analytics over tenders, bids, suppliers, and invoices",
    "Procurement Integrity Analyst",
    "a watchdog over every tender — answers with numbers and charts, not opinions",
    [
        ("Supplier Risk Score",
         "One score per supplier,\nbuilt from eight risk signals."),
        ("Bid-Rigging Screen",
         "Rotation rings and cover\nbids across related firms."),
        ("Price Anomalies",
         "Overpricing vs the market\nbenchmark, quantified."),
        ("Threshold & Invoice Checks",
         "Split purchases under limits,\nduplicate invoices."),
    ],
    ["43A87C", "2B9166", "15794F", "0C5F3C"],
    "Governed Procurement Data  ·  Tenders · Bids · Suppliers · Invoices",
    "Every figure the agent quotes is traceable to a query you can inspect — drill into any number in chat",
)

# ── 3 · Global Intelligence (amber) ─────────────────────────────────
agent_slide(
    "GLOBAL INTELLIGENCE  ·  EYES ON THE WORLD",
    "What other governments are doing — and what it means for NCGR",
    "Global Intelligence Agent",
    "scans news and publications worldwide — every answer arrives with its sources",
    [
        ("Country Benchmarks",
         "How peers run AI-driven\noversight — Korea, Brazil, EU…"),
        ("Emerging Technology",
         "What's new, summarized —\nand what it means for you."),
        ("Topic Monitoring",
         "A standing watch on the\nthemes you care about."),
        ("Cited Sources",
         "Open the original article\nbehind any claim, one click."),
    ],
    ["D99A2B", "C08321", "A66C17", "8A5710"],
    "Live Web & News Search",
    "Grounded and current — the agent quotes what it read, links it, and never invents a source",
)

prs.core_properties.title = "NCGR Agent Slides — Investigation · Procurement · Global Intelligence"
OUT = str(Path(__file__).resolve().parent / "NCGR_Agent_Slides.pptx")
prs.save(OUT)
print("saved", OUT)
